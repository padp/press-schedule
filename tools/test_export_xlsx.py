"""Tests for export_xlsx.py, run entirely against tools/test_output/
(gitignored) - never the real network share. See CLAUDE.md's "Do not
overwrite real files while building/testing this".

Seeded with the same real Press 2 data used elsewhere in this project (the
actual rows read from D-3 Press Report - Thursday 2nd Shift.xls - see
CLAUDE.md), including a numeric-typed column (oven_cavity) and a note row,
so this proves the export handles the real shape, not an invented one.

    pip install --user pytest openpyxl
    python -m pytest tools\\test_export_xlsx.py -v
"""
import datetime
import os

import openpyxl
import pytest

from export_xlsx import write_schedule_xlsx

OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "test_output")

REAL_PRESS2_DOC = {
    "press": "PRESS 2",
    "day_of_week": "Thursday",
    "shift": "2nd",
    "date": "2026-09-24",
    "roster": {
        "press_op": "Dylan", "saw_op_1": "Ryan", "saw_op_2": "Will",
        "oven_probes": "DG - 3:35 PM", "supervisor": "Wally G",
    },
    "columns": ["die_no", "suffix", "job_no", "part_no", "alloy_temper",
                "blts", "cut_length", "est_wt_ft", "cast_no", "blt_length",
                "blts_ran", "die_temp", "oven_cavity", "start_time",
                "stop_time", "die_failure", "die_pull", "downtime_code",
                "time_down", "comments"],
    "column_types": {"oven_cavity": "number", "time_down": "number"},
    "rows": [
        {"kind": "job", "die_no": "1309", "suffix": "", "job_no": "25892",
         "part_no": "1309X-1", "alloy_temper": "6005AT6", "blts": "43",
         "cut_length": "232.0", "est_wt_ft": "1.056", "oven_cavity": "3",
         "start_time": "06:44", "stop_time": "07:14", "time_down": "5.5"},
        # real entry: "BAL." in a column that's NOT numeric-typed - must
        # stay a plain string in the exported file, not be coerced
        {"kind": "job", "die_no": "1307", "job_no": "25893",
         "part_no": "1307X-2", "alloy_temper": "6005AT6", "blts": "BAL.",
         "cut_length": "202.9", "est_wt_ft": "1.716"},
        {"kind": "note", "text": "ALLOY CHANGE to"},
    ],
}


@pytest.fixture(autouse=True)
def clean_output():
    os.makedirs(OUT_DIR, exist_ok=True)
    yield
    for f in os.listdir(OUT_DIR):
        os.remove(os.path.join(OUT_DIR, f))


def _out(name):
    return os.path.join(OUT_DIR, name)


def test_writes_a_real_readable_xlsx_file():
    path = write_schedule_xlsx(REAL_PRESS2_DOC, _out("basic.xlsx"))
    assert os.path.exists(path)
    wb = openpyxl.load_workbook(path)
    assert wb.active is not None


def test_title_and_date_in_the_header():
    wb = openpyxl.load_workbook(write_schedule_xlsx(REAL_PRESS2_DOC, _out("t.xlsx")))
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "PRESS 2"
    # date cell is the last column of row 1
    assert ws.cell(row=1, column=ws.max_column).value == "2026-09-24"


def test_roster_names_present():
    wb = openpyxl.load_workbook(write_schedule_xlsx(REAL_PRESS2_DOC, _out("r.xlsx")))
    ws = wb.active
    values = [c.value for row in ws.iter_rows() for c in row]
    for name in ("Dylan", "Ryan", "Will", "DG - 3:35 PM", "Wally G"):
        assert name in values, f"{name!r} missing from exported roster"


def test_column_headers_match_labels():
    wb = openpyxl.load_workbook(write_schedule_xlsx(REAL_PRESS2_DOC, _out("h.xlsx")))
    ws = wb.active
    values = [c.value for row in ws.iter_rows() for c in row]
    assert "Die #" in values
    assert "Downtime Code" in values
    # numeric columns are flagged in their own header text
    assert any(v and "Oven Cavity" in str(v) and "#" in str(v) for v in values)


def test_numeric_column_is_a_real_excel_number_not_text():
    """The whole point of tracking column_types through to the export -
    oven_cavity should be openpyxl type int/float, not the string "3"."""
    wb = openpyxl.load_workbook(write_schedule_xlsx(REAL_PRESS2_DOC, _out("n.xlsx")))
    ws = wb.active
    oven_cavity_values = [
        c.value for row in ws.iter_rows() for c in row if c.value == 3
    ]
    assert oven_cavity_values, "expected a real numeric 3, not the string '3'"
    assert isinstance(oven_cavity_values[0], (int, float))

    time_down_values = [
        c.value for row in ws.iter_rows() for c in row if c.value == 5.5
    ]
    assert time_down_values and isinstance(time_down_values[0], float)


def test_bal_stays_a_string_in_a_non_numeric_column():
    """The other job row's "blts": "BAL." must NOT be coerced to a number
    just because it happens to sit next to numeric columns."""
    wb = openpyxl.load_workbook(write_schedule_xlsx(REAL_PRESS2_DOC, _out("b.xlsx")))
    ws = wb.active
    values = [c.value for row in ws.iter_rows() for c in row]
    assert "BAL." in values


def test_time_columns_become_real_excel_times():
    wb = openpyxl.load_workbook(write_schedule_xlsx(REAL_PRESS2_DOC, _out("time.xlsx")))
    ws = wb.active
    time_values = [c.value for row in ws.iter_rows() for c in row
                   if isinstance(c.value, datetime.time)]
    assert datetime.time(6, 44) in time_values
    assert datetime.time(7, 14) in time_values


def test_note_row_present_and_distinct_from_job_rows():
    wb = openpyxl.load_workbook(write_schedule_xlsx(REAL_PRESS2_DOC, _out("note.xlsx")))
    ws = wb.active
    values = [c.value for row in ws.iter_rows() for c in row]
    assert "ALLOY CHANGE to" in values


def test_missing_roster_or_columns_does_not_crash():
    """A brand-new, never-filled-in slot shouldn't make the export blow up."""
    sparse = {"press": "PRESS 2", "day_of_week": "Sunday", "shift": "1st"}
    path = write_schedule_xlsx(sparse, _out("sparse.xlsx"))
    assert os.path.exists(path)


def test_output_never_touches_anything_outside_test_output():
    """Guard against a future edit accidentally hardcoding a real path -
    every test in this file must write under OUT_DIR only."""
    before = set(os.listdir(OUT_DIR))
    write_schedule_xlsx(REAL_PRESS2_DOC, _out("guard.xlsx"))
    after = set(os.listdir(OUT_DIR))
    assert after - before == {"guard.xlsx"}
