"""Writes one schedule document out as a real .xlsx file - the one-way,
non-authoritative snapshot described in README.md's Architecture section.

This is deliberately NOT a byte-for-byte reproduction of the real Press 2
sheet's exact row/column pixel positions (that sheet pads its job grid with
~55 blank rows before the roster panel, matched to a printed paper form -
see CLAUDE.md). Nobody scrolling a generated snapshot benefits from
reproducing that; what matters is the same INFORMATION, laid out compactly:
title + date, the roster panel, then the job/note grid in real run order.

Column typing is preserved, not flattened to text: a column listed in
column_types as "number" is written as a real Excel number (openpyxl
infers the cell type from the Python value), not a text string - so a
generated report is at least as useful to open in real Excel as the
current process's files are, not a regression.

Usage:
    from export_xlsx import write_schedule_xlsx
    write_schedule_xlsx(doc, r"C:\\some\\path\\PRESS 2 Thursday 2nd.xlsx")

CAUTION - see CLAUDE.md's "Do not overwrite real files while building/
testing this": never call this against a real path under Press Schedules\\
or Press Reports\\ until the scheduled export job that will eventually own
that responsibility is built and reviewed. Point it at tools/test_output/
(gitignored) for all development and testing.
"""
import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TITLE_FILL = PatternFill("solid", fgColor="D9E2F3")
ROSTER_FILL = PatternFill("solid", fgColor="F2F2F2")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(color="FFFFFF", bold=True)
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE_FONT = Font(italic=True, bold=True)

ROSTER_LABELS = {
    "press_op": "Press Op", "saw_op_1": "Saw Op #1", "saw_op_2": "Saw Op #2",
    "oven_probes": "Oven Probes", "supervisor": "Supervisor",
    "reviewed_by": "Reviewed By",
}

COLUMN_LABELS = {
    "die_no": "Die #", "suffix": "Suffix", "job_no": "Job #",
    "part_no": "Part #", "alloy_temper": "Alloy/temper", "blts": "# blts",
    "cut_length": "Cut length", "est_wt_ft": "Est wt/ft", "cast_no": "Cast #",
    "blt_length": "blt length", "blts_ran": "blts ran", "die_temp": "Die Temp",
    "oven_cavity": "Oven Cavity", "start_time": "Start time",
    "stop_time": "Stop time", "die_failure": "Die Failure",
    "die_pull": "Die Pull", "downtime_code": "Downtime Code",
    "time_down": "Time Down (min)", "comments": "Comments",
    "str_blt_length": "Str blt length",
}

TIME_COLUMNS = {"start_time", "stop_time"}


def _column_label(key):
    return COLUMN_LABELS.get(key, key.replace("_", " ").title())


def _cell_value(col, value, column_types):
    """Numbers stay numbers, times become real Excel times, everything
    else (including a deliberately-non-numeric entry like "BAL." in a
    column NOT listed as numeric) stays a plain string - see README's
    "Deliberate choices" for why free text is the default, not typed.
    """
    if value in (None, ""):
        return None
    if col in TIME_COLUMNS and isinstance(value, str) and ":" in value:
        try:
            h, m = value.split(":")
            return datetime.time(int(h), int(m))
        except (ValueError, TypeError):
            return value  # not a clean HH:MM - leave as the original text
    if (column_types or {}).get(col) == "number":
        try:
            f = float(value)
            return int(f) if f.is_integer() else f
        except (TypeError, ValueError):
            return value  # shouldn't happen (API validates this), but
                          # never let export raise over bad input either
    return value


def write_schedule_xlsx(doc, path):
    """doc: a schedule document shaped like README's data model (press,
    day_of_week, shift, date, roster, columns, column_types, rows).
    path: full output file path - caller's responsibility to point this at
    a safe location (see module docstring)."""
    wb = Workbook()
    ws = wb.active
    ws.title = f"{doc.get('day_of_week', '')} {doc.get('shift', '')}"[:31] or "Schedule"

    columns = doc.get("columns") or []
    column_types = doc.get("column_types") or {}
    n_cols = max(len(columns), 1)

    # --- title + date -------------------------------------------------
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(n_cols - 2, 1))
    title_cell = ws.cell(row=1, column=1, value=doc.get("press", ""))
    title_cell.font = Font(bold=True, size=14)
    title_cell.fill = TITLE_FILL
    date_label = ws.cell(row=1, column=max(n_cols - 1, 2), value="Date:")
    date_label.font = Font(bold=True)
    date_cell = ws.cell(row=1, column=n_cols, value=doc.get("date") or "")

    # --- roster panel ---------------------------------------------------
    roster = doc.get("roster") or {}
    row = 3
    for key, value in roster.items():
        label = ROSTER_LABELS.get(key, key.replace("_", " ").title())
        lc = ws.cell(row=row, column=1, value=f"{label}:")
        lc.font = Font(bold=True)
        lc.fill = ROSTER_FILL
        ws.cell(row=row, column=2, value=value or "")
        row += 1
    row += 1  # blank spacer row before the grid

    # --- column headers ---------------------------------------------------
    header_row = row
    for i, col in enumerate(columns, start=1):
        c = ws.cell(row=header_row, column=i, value=_column_label(col))
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        if column_types.get(col) == "number":
            c.value = f"{_column_label(col)} (#)"
    row += 1

    # --- rows: job rows get one cell per column, note rows span the width -
    for r in doc.get("rows") or []:
        if r.get("kind") == "note":
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
            c = ws.cell(row=row, column=1, value=r.get("text", ""))
            c.font = NOTE_FONT
            c.fill = NOTE_FILL
        else:
            for i, col in enumerate(columns, start=1):
                ws.cell(row=row, column=i, value=_cell_value(col, r.get(col), column_types))
        row += 1

    for i in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14

    wb.save(path)
    return path
